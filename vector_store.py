"""
Vector Store Module for RAG Application
Handles document processing, chunking, embeddings, and FAISS index management.
"""

import os
import json
import uuid
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime

import faiss
import numpy as np
from sentence_transformers import SentenceTransformer

# Document processing libraries
try:
    import fitz  # PyMuPDF for PDF
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class VectorStore:
    """Handles document chunking, embedding, and retrieval using FAISS."""
    
    def __init__(
        self,
        chunk_size: int = 1000,
        chunk_overlap: int = 100,
        embedding_model: str = "all-MiniLM-L6-v2",
        top_k: int = 3
    ):
        """
        Initialize the vector store.
        
        Args:
            chunk_size: Size of each text chunk in characters
            chunk_overlap: Overlap between chunks in characters
            embedding_model: Sentence transformer model name
            top_k: Number of chunks to retrieve per query
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.top_k = top_k
        
        # Load embedding model
        logger.info(f"Loading embedding model: {embedding_model}")
        self.embedding_model = SentenceTransformer(embedding_model)
        self.embedding_dim = self.embedding_model.get_sentence_embedding_dimension()
        
        # Initialize FAISS index
        self.index = faiss.IndexFlatIP(self.embedding_dim)
        
        # Store chunks with metadata
        self.chunks: List[Dict[str, Any]] = []
        self.documents: Dict[str, Dict[str, Any]] = {}
        
        logger.info(f"Vector store initialized: chunk_size={chunk_size}, overlap={chunk_overlap}, top_k={top_k}")
    
    def _extract_text_from_pdf(self, file_path: Path) -> str:
        """Extract text from PDF file."""
        if not PDF_AVAILABLE:
            raise ImportError("PyMuPDF not installed. Install with: pip install pymupdf")
        
        text_parts = []
        doc = fitz.open(file_path)
        
        for page_num, page in enumerate(doc):
            page_text = page.get_text()
            if page_text.strip():
                text_parts.append(page_text)
        
        doc.close()
        return "\n\n".join(text_parts)
    
    def _extract_text_from_docx(self, file_path: Path) -> str:
        """Extract text from Word document."""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx not installed. Install with: pip install python-docx")
        
        doc = docx.Document(file_path)
        text_parts = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # Extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        
        return "\n\n".join(text_parts)
    
    def _extract_text_from_xlsx(self, file_path: Path) -> str:
        """Extract text from Excel file."""
        if not XLSX_AVAILABLE:
            raise ImportError("openpyxl not installed. Install with: pip install openpyxl")
        
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text_parts = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"Sheet: {sheet_name}")
            
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text_parts.append(row_text)
        
        return "\n".join(text_parts)
    
    def _extract_text_from_pptx(self, file_path: Path) -> str:
        """Extract text from PowerPoint file."""
        if not PPTX_AVAILABLE:
            raise ImportError("python-pptx not installed. Install with: pip install python-pptx")
        
        prs = Presentation(file_path)
        text_parts = []
        
        for slide_num, slide in enumerate(prs.slides):
            text_parts.append(f"Slide {slide_num + 1}:")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text)
        
        return "\n\n".join(text_parts)
    
    def _extract_text_from_file(self, file_path: Path) -> str:
        """Extract text from various file types."""
        suffix = file_path.suffix.lower()
        
        if suffix == ".pdf":
            return self._extract_text_from_pdf(file_path)
        elif suffix in [".docx", ".doc"]:
            return self._extract_text_from_docx(file_path)
        elif suffix in [".xlsx", ".xls"]:
            return self._extract_text_from_xlsx(file_path)
        elif suffix in [".pptx", ".ppt"]:
            return self._extract_text_from_pptx(file_path)
        elif suffix in [".txt", ".md", ".markdown"]:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        elif suffix == ".csv":
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        else:
            raise ValueError(f"Unsupported file type: {suffix}")
    
    def _create_chunks(self, text: str, filename: str) -> List[Dict[str, Any]]:
        """
        Split text into overlapping chunks with metadata.
        
        Args:
            text: Full document text
            filename: Source filename for metadata
            
        Returns:
            List of chunk dictionaries with text and metadata
        """
        chunks = []
        
        # Clean text
        text = text.strip()
        if not text:
            return chunks
        
        # Create chunks with overlap
        start = 0
        chunk_index = 0
        
        while start < len(text):
            end = start + self.chunk_size
            chunk_text = text[start:end]
            
            # Try to break at sentence boundary if possible
            if end < len(text):
                last_period = chunk_text.rfind(". ")
                last_newline = chunk_text.rfind("\n")
                break_point = max(last_period, last_newline)
                
                if break_point > self.chunk_size * 0.5:
                    chunk_text = chunk_text[:break_point + 1]
                    end = start + break_point + 1
            
            if chunk_text.strip():
                chunk_id = f"chunk_{filename}_{chunk_index}"
                
                chunks.append({
                    "id": chunk_id,
                    "text": chunk_text.strip(),
                    "filename": filename,
                    "chunk_index": chunk_index,
                    "start_char": start,
                    "end_char": end
                })
                chunk_index += 1
            
            # Move to next chunk with overlap
            start = end - self.chunk_overlap
            if start <= 0 and end >= len(text):
                break
            start = max(start, end - self.chunk_overlap) if end < len(text) else len(text)
        
        return chunks
    
    def add_document(self, file_path: str) -> Dict[str, Any]:
        """
        Process and add a document to the vector store.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Document metadata dictionary
        """
        file_path = Path(file_path)
        filename = file_path.name
        
        logger.info(f"Processing document: {filename}")
        
        # Extract text
        text = self._extract_text_from_file(file_path)
        
        if not text.strip():
            raise ValueError(f"No text extracted from {filename}")
        
        # Create chunks
        doc_chunks = self._create_chunks(text, filename)
        
        if not doc_chunks:
            raise ValueError(f"No chunks created from {filename}")
        
        # Generate embeddings
        chunk_texts = [c["text"] for c in doc_chunks]
        embeddings = self.embedding_model.encode(chunk_texts, normalize_embeddings=True)
        
        # Add to FAISS index
        start_idx = len(self.chunks)
        self.index.add(embeddings.astype(np.float32))
        
        # Store chunks with their index positions
        for i, chunk in enumerate(doc_chunks):
            chunk["index_position"] = start_idx + i
            self.chunks.append(chunk)
        
        # Store document metadata
        doc_id = f"doc_{uuid.uuid4().hex[:8]}"
        self.documents[doc_id] = {
            "document_id": doc_id,
            "filename": filename,
            "file_path": str(file_path),
            "chunk_count": len(doc_chunks),
            "total_characters": len(text),
            "added_at": datetime.now().isoformat()
        }
        
        logger.info(f"Added {len(doc_chunks)} chunks from {filename}")
        
        return self.documents[doc_id]
    
    def add_documents(self, file_paths: List[str]) -> List[Dict[str, Any]]:
        """
        Process and add multiple documents to the vector store.
        
        Args:
            file_paths: List of paths to document files
            
        Returns:
            List of document metadata dictionaries
        """
        results = []
        
        for file_path in file_paths:
            try:
                doc_meta = self.add_document(file_path)
                results.append(doc_meta)
            except Exception as e:
                logger.error(f"Error processing {file_path}: {str(e)}")
                results.append({
                    "filename": Path(file_path).name,
                    "error": str(e)
                })
        
        return results
    
    def search(self, query: str, top_k: Optional[int] = None) -> List[Dict[str, Any]]:
        """
        Search for relevant chunks based on query.
        
        Args:
            query: Search query
            top_k: Number of results (default: self.top_k)
            
        Returns:
            List of matching chunks with similarity scores
        """
        if self.index.ntotal == 0:
            return []
        
        top_k = top_k or self.top_k
        top_k = min(top_k, self.index.ntotal)
        
        # Embed query
        query_embedding = self.embedding_model.encode([query], normalize_embeddings=True)
        
        # Search FAISS index
        scores, indices = self.index.search(query_embedding.astype(np.float32), top_k)
        
        # Build results with citation metadata
        results = []
        for score, idx in zip(scores[0], indices[0]):
            if idx < len(self.chunks):
                chunk = self.chunks[idx].copy()
                chunk["similarity"] = float(score)
                chunk["cite_as"] = f"[{chunk['filename']}](chunk_{chunk['chunk_index']})"
                results.append(chunk)
        
        return results
    
    def get_context_for_llm(self, query: str, top_k: Optional[int] = None) -> Tuple[str, List[Dict[str, Any]]]:
        """
        Get formatted context string for LLM from search results.
        
        Args:
            query: User query
            top_k: Number of chunks to retrieve
            
        Returns:
            Tuple of (formatted context string, list of source chunks)
        """
        chunks = self.search(query, top_k)
        
        if not chunks:
            return "", []
        
        context_parts = []
        
        for i, chunk in enumerate(chunks):
            context_parts.append(f"[Excerpt {i+1} from {chunk['filename']}]")
            context_parts.append(f"{chunk['text']}")
            context_parts.append("-" * 40)
        
        return "\n".join(context_parts), chunks
    
    def get_stats(self) -> Dict[str, Any]:
        """Get vector store statistics."""
        return {
            "total_documents": len(self.documents),
            "total_chunks": len(self.chunks),
            "index_size": self.index.ntotal,
            "embedding_dimension": self.embedding_dim,
            "chunk_size": self.chunk_size,
            "chunk_overlap": self.chunk_overlap,
            "documents": list(self.documents.values())
        }
    
    def clear(self):
        """Clear all documents and reset the index."""
        self.index = faiss.IndexFlatIP(self.embedding_dim)
        self.chunks = []
        self.documents = {}
        logger.info("Vector store cleared")